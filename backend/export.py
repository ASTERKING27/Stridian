"""
Excel downloads: a whole squad (every sport, for an admin) or a single student.

Built from the same rows as the Google Sheets, so a download and the admin's sheets
always agree. Coaches get these instead of access to the sheets themselves.
"""

import re
from io import BytesIO

import sheets
import sports_config as sc

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# control characters make Excel refuse the whole file (the same set openpyxl rejects)
ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")

RESULT_HEADER = ["Student ID", "RA number", "Student", "Sport", "Test", "Value", "Unit",
                 "Recorded on"]


def _clean(value):
    return ILLEGAL.sub("", value) if isinstance(value, str) else value


def _workbook():
    # imported here so the worker computer, which never makes downloads, doesn't need it
    from openpyxl import Workbook
    return Workbook()


def _tab(wb, title, header, rows):
    from openpyxl.styles import Font

    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append([_clean(v) for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for column in ws.iter_cols():
        for cell in column:
            # anything a person typed stays text: openpyxl would make "=..." a live formula
            if cell.data_type == "f":
                cell.data_type = "s"
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(10, longest + 2), 60)
    return ws


def _result_rows(student, history):
    metrics = sc.metric_map(student.sport)
    return [
        [student.id, student.ra_number or "", student.name, student.sport,
         metrics.get(h["metric_key"], {}).get("label", h["metric_key"]), h["value"],
         metrics.get(h["metric_key"], {}).get("unit", ""), sheets._local(h["recorded_at"])]
        for h in history
    ]


def _bytes(wb) -> bytes:
    del wb[wb.sheetnames[0]]     # the empty sheet every new workbook starts with
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def squad(students, histories: dict, coaches=None) -> bytes:
    """Every student given, one tab per kind of information. `histories` maps a
    student id to their test results, oldest first."""
    wb = _workbook()
    _tab(wb, "Profiles", sheets.PROFILE_HEADER, [sheets.profile_row(s) for s in students])
    _tab(wb, "Analysis", sheets.HEADER,
         [s.sheet_row for s in students if s.sheet_row and len(s.sheet_row) == len(sheets.HEADER)])
    _tab(wb, "Achievements", sheets.ACHIEVEMENT_HEADER,
         [sheets.achievement_row(a) for s in students for a in s.achievements if a.status != "draft"])
    _tab(wb, "Test results", RESULT_HEADER,
         [row for s in students for row in _result_rows(s, histories.get(s.id, []))])
    if coaches is not None:
        _tab(wb, "Coaches", sheets.COACH_HEADER, [sheets.coach_row(c) for c in coaches])
    return _bytes(wb)


def student(s, report: dict, history: list) -> bytes:
    """One student: their record, where they fit, every metric, achievements, history."""
    wb = _workbook()
    _tab(wb, "Profile", ["Field", "Value"], zip(sheets.PROFILE_HEADER, sheets.profile_row(s)))
    _tab(wb, "Positions", ["Position", "Fit /100", "Confidence", "Coverage"],
         [[p["position"], p["fit"], p["confidence"], p["coverage"]] for p in report["positions"]])
    _tab(wb, "Metrics", ["Metric", "Value", "Unit", "Score /100", "From"],
         [[m["label"], m["value"], m["unit"], m["score"], m["source"]] for m in report["metrics"]])
    _tab(wb, "Achievements", sheets.ACHIEVEMENT_HEADER,
         [sheets.achievement_row(a) for a in s.achievements if a.status != "draft"])
    _tab(wb, "Test results", RESULT_HEADER, _result_rows(s, history))
    return _bytes(wb)
